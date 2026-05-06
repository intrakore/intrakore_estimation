# intrakore_estimation/api.py
import frappe
from frappe import _
from frappe.utils import getdate, today
import json

import frappe
import json

@frappe.whitelist()
def create_bid(bid_data):
    print("create_bid called with data:", bid_data)  # Debug log
    """Create a new bid document - handles customer creation/linking"""
    try:
        if isinstance(bid_data, str):
            bid_data = json.loads(bid_data)
        
        frappe.log_error(f"create_bid called with data: {bid_data}", "Intrakore Estimation")
        
        # Validate required fields
        if not bid_data.get("bid_code"):
            return {"error": "bid_code is required"}
        
        if not bid_data.get("project_name"):
            return {"error": "project_name is required"}
        
        # Check if bid already exists
        if frappe.db.exists("Bid", {"bid_code": bid_data["bid_code"]}):
            bid = frappe.get_doc("Bid", {"bid_code": bid_data["bid_code"]})
            return bid.as_dict()
        
        # Handle client/customer
        client_name = None
        client_link = None
        
        if bid_data.get("bid_source") == "CRM-linked" and bid_data.get("opportunity"):
            # For CRM-linked, get client from opportunity
            opportunity = frappe.get_doc("Opportunity", bid_data.get("opportunity"))
            client_link = opportunity.client
            client_name = opportunity.client_name
            
        elif bid_data.get("client_name"):
            # For standalone, try to find or create customer
            client_name = bid_data.get("client_name")
            
            # Check if customer exists by name
            existing_customer = frappe.db.get_value("Customer", {"customer_name": client_name}, "name")
            
            if existing_customer:
                client_link = existing_customer
                frappe.log_error(f"Found existing customer: {client_link}", "Intrakore Estimation")
            else:
                # Create a new customer
                try:
                    new_customer = frappe.get_doc({
                        "doctype": "Customer",
                        "customer_name": client_name,
                        "customer_group": "Commercial",
                        "customer_type": "Company",
                        "territory": "United Arab Emirates"
                    })
                    new_customer.insert(ignore_permissions=True)
                    client_link = new_customer.name
                    frappe.log_error(f"Created new customer: {client_link}", "Intrakore Estimation")
                except Exception as e:
                    frappe.log_error(f"Error creating customer: {str(e)}", "Intrakore Estimation")
                    # If customer creation fails, still proceed with just the name
                    client_link = None
        
        # Also handle client_name if provided directly
        if not client_name and bid_data.get("client_name"):
            client_name = bid_data.get("client_name")
        
        # Prepare bid dictionary
        bid_dict = {
            "doctype": "Bid",
            "bid_code": bid_data["bid_code"],
            "project_name": bid_data["project_name"],
            "bid_source": bid_data.get("bid_source", "Standalone"),
            "status": bid_data.get("status", "Draft"),
            "internal_notes": bid_data.get("internal_notes", "")
        }
        
        # Add client link (if we have one)
        if client_link:
            bid_dict["client"] = client_link
        
        # Add client name (always have this)
        if client_name:
            bid_dict["client_name"] = client_name
        elif bid_data.get("client_name"):
            bid_dict["client_name"] = bid_data.get("client_name")
        
        # Add opportunity if provided
        if bid_data.get("opportunity"):
            bid_dict["opportunity"] = bid_data.get("opportunity")
        
        # Add lead estimator if provided
        if bid_data.get("lead_estimator"):
            bid_dict["lead_estimator"] = bid_data.get("lead_estimator")
            # Fetch the full name
            user_info = frappe.db.get_value("User", bid_data["lead_estimator"], "full_name")
            if user_info:
                bid_dict["lead_estimator_name"] = user_info
        
        # Add submission date if provided
        if bid_data.get("submission_date"):
            bid_dict["submission_date"] = bid_data.get("submission_date")
        
        frappe.log_error(f"Creating bid with dict: {bid_dict}", "Intrakore Estimation")
        
        # Create the bid
        bid = frappe.get_doc(bid_dict)
        bid.insert(ignore_permissions=True)
        frappe.db.commit()
        
        return bid.as_dict()
        
    except Exception as e:
        frappe.log_error(f"Error in create_bid: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}

@frappe.whitelist()
def get_pricing_data(bid_id):
    print(f"get_pricing_data called with bid_id: {bid_id}")  # Debug log
    """Get pricing data for a bid"""
    try:
        # Get the bid
        bid = frappe.get_doc("Bid", bid_id)
        
        # Get BOQ lines (you'll need to create this child table)
        # For now, return sample structure
        packages = [
            {"name": "Preliminaries", "color": "blue", "total": 14, "confirmed": 0},
            {"name": "Substructure", "color": "green", "total": 22, "confirmed": 0},
            {"name": "Superstructure", "color": "purple", "total": 31, "confirmed": 0},
        ]
        
        lines = [
            {"id": 1, "package": "Superstructure", "item": "3.1.1", "description": "Reinforced concrete columns", "qty": 186, "unit": "m³", "rate": 2540, "amount": 472440, "confirmed": False},
            {"id": 2, "package": "Superstructure", "item": "3.1.2", "description": "Reinforced concrete beams", "qty": 142, "unit": "m³", "rate": 2378, "amount": 337676, "confirmed": True},
        ]
        
        return {"packages": packages, "lines": lines}
        print(result)  # Debug log of the result
    except Exception as e:
        frappe.log_error(f"Error in get_pricing_data: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}


@frappe.whitelist()
def save_pricing_data(bid_id, lines, packages):
    """Save pricing data for a bid"""
    try:
        if isinstance(lines, str):
            lines = json.loads(lines)
        if isinstance(packages, str):
            packages = json.loads(packages)
        
        # Save to custom fields or child table
        frappe.db.set_value("Bid", bid_id, "custom_pricing_data", json.dumps(lines))
        frappe.db.commit()
        
        return {"success": True}
        
    except Exception as e:
        frappe.log_error(f"Error in save_pricing_data: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}
        
@frappe.whitelist()
def save_package_tagging(bid_id, lines):
    print(f"save_package_tagging called with bid_id: {bid_id} and lines: {lines}")  # Debug log
    """Save package assignments for BOQ lines"""
    try:
        if isinstance(lines, str):
            lines = json.loads(lines)
        
        # Get the bid document
        bid = frappe.get_doc("Bid", bid_id)
        
        # Store tagging data (you may want to create a child table for this)
        bid.db_set("custom_package_data", json.dumps(lines))
        
        frappe.db.commit()
        return {"success": True, "message": "Package tagging saved"}
        
    except Exception as e:
        frappe.log_error(f"Error in save_package_tagging: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}


@frappe.whitelist()
def save_pricing(bid_id, lines, packages):
    """Save pricing data for BOQ lines"""
    try:
        if isinstance(lines, str):
            lines = json.loads(lines)
        if isinstance(packages, str):
            packages = json.loads(packages)
        
        bid = frappe.get_doc("Bid", bid_id)
        
        # Store pricing data
        bid.db_set("custom_pricing_data", json.dumps(lines))
        bid.db_set("total_priced_value", sum(line.get("amount", 0) for line in lines))
        
        frappe.db.commit()
        return {"success": True, "message": "Pricing saved"}
        
    except Exception as e:
        frappe.log_error(f"Error in save_pricing: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}


@frappe.whitelist()
def save_bid_strategy(bid_id, strategy_data):
    print(f"save_bid_strategy called with bid_id: {bid_id} and strategy_data: {strategy_data}")  # Debug log
    """Save bid strategy data"""
    try:
        if isinstance(strategy_data, str):
            strategy_data = json.loads(strategy_data)
        
        bid = frappe.get_doc("Bid", bid_id)
        
        # Store strategy data
        bid.db_set("custom_strategy_data", json.dumps(strategy_data))
        bid.db_set("commission_rate", strategy_data.get("acquisitionPercent", 0))
        bid.db_set("bid_total", strategy_data.get("bidTotalWithAcquisition", 0))
        
        frappe.db.commit()
        return {"success": True, "message": "Bid strategy saved"}
        
    except Exception as e:
        frappe.log_error(f"Error in save_bid_strategy: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}


@frappe.whitelist()
def submit_bid_for_review(submit_data):
    print(f"submit_bid_for_review called with submit_data: {submit_data}")  # Debug log
    """Submit bid for CM review"""
    try:
        if isinstance(submit_data, str):
            submit_data = json.loads(submit_data)
        
        bid_code = submit_data.get("bid_code")
        bid = frappe.get_doc("Bid", {"bid_code": bid_code})
        
        # Update status
        bid.db_set("status", "Review")
        bid.db_set("bid_total", submit_data.get("bid_total", 0))
        
        frappe.db.commit()
        return {"success": True, "message": f"Bid {bid_code} submitted for review"}
        
    except Exception as e:
        frappe.log_error(f"Error in submit_bid_for_review: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}

@frappe.whitelist()
def get_opportunities(filters=None, search=None):
    """Fetch list of open opportunities"""
    try:
        conditions = []
        values = {}
        
        # Filter by stage - exclude closed won/lost
        conditions.append("stage NOT IN ('Closed Won', 'Closed Lost')")
        
        if search:
            conditions.append("(opportunity_name LIKE %(search)s OR client_name LIKE %(search)s)")
            values["search"] = f"%{search}%"
        
        where_clause = " AND ".join(conditions)
        
        opportunities = frappe.db.sql(f"""
            SELECT 
                name,
                opportunity_name,
                client,
                client_name,
                estimated_value,
                probability,
                stage,
                expected_close_date
            FROM `tabOpportunity`
            WHERE {where_clause}
            ORDER BY expected_close_date ASC
            LIMIT 50
        """, values, as_dict=True)
        
        # Format for dropdown
        formatted = [{
            "value": opp.name,
            "label": f"{opp.opportunity_name} · {opp.client_name or opp.client} · AED {opp.estimated_value:,.0f}",
            "data": opp
        } for opp in opportunities]
        
        return {"opportunities": formatted, "total": len(formatted)}
        
    except Exception as e:
        frappe.log_error(f"Error in get_opportunities: {str(e)}", "Intrakore Estimation")
        return {"opportunities": [], "total": 0, "error": str(e)}

@frappe.whitelist()
def get_users_by_role(roles=None, txt=None, page=1, page_size=20):
    """
    Fetch users with specific roles (Estimator, Estimation Manager, etc.)
    
    Args:
        roles: Comma-separated list of role names
        txt: Search text for user name/email
        page: Page number for pagination
        page_size: Number of records per page
    """
    try:
        if isinstance(roles, str):
            roles = roles.split(',') if roles else []
        
        if not roles:
            roles = ['Estimator', 'Estimation Manager']
        
        # Build role condition
        role_conditions = []
        values = {}
        
        for i, role in enumerate(roles):
            role_conditions.append(f"hr.role = %(role_{i})s")
            values[f"role_{i}"] = role.strip()
        
        role_clause = " OR ".join(role_conditions)
        
        # Calculate offset
        page = int(page) if page else 1
        page_size = int(page_size) if page_size else 20
        offset = (page - 1) * page_size
        
        # Build search condition
        search_condition = ""
        if txt:
            values["search"] = f"%{txt}%"
            search_condition = "AND (u.name LIKE %(search)s OR u.full_name LIKE %(search)s)"
        
        # Query users with specific roles
        query = f"""
            SELECT DISTINCT 
                u.name,
                u.full_name,
                u.first_name,
                u.last_name,
                u.user_image,
                u.enabled
            FROM `tabUser` u
            INNER JOIN `tabHas Role` hr ON hr.parent = u.name
            WHERE u.enabled = 1
            AND ({role_clause})
            {search_condition}
            ORDER BY u.full_name ASC
            LIMIT %(page_size)s OFFSET %(offset)s
        """
        
        values["page_size"] = page_size
        values["offset"] = offset
        
        users = frappe.db.sql(query, values, as_dict=True)
        
        # Get total count for pagination
        count_query = f"""
            SELECT COUNT(DISTINCT u.name) as total
            FROM `tabUser` u
            INNER JOIN `tabHas Role` hr ON hr.parent = u.name
            WHERE u.enabled = 1
            AND ({role_clause})
            {search_condition}
        """
        
        total_result = frappe.db.sql(count_query, values, as_dict=True)
        total = total_result[0].total if total_result else 0
        
        return {
            "users": users,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size
        }
        
    except Exception as e:
        frappe.log_error(f"Error in get_users_by_role: {str(e)}", "Intrakore Estimation")
        return {"users": [], "total": 0, "error": str(e)}


@frappe.whitelist()
def get_lead_estimators(txt=None, page=1, page_size=20):
    """Convenience method to get users with Estimator or Estimation Manager roles"""
    return get_users_by_role(roles="Estimator,Estimation Manager", txt=txt, page=page, page_size=page_size)
    
@frappe.whitelist()
def create_new_bid():
    """Create a new blank bid"""
    try:
        from datetime import datetime
        import random
        
        year = datetime.now().strftime('%y')
        month = datetime.now().strftime('%m')
        random_num = random.randint(100, 999)
        bid_code = f"INT-{year}{month}-{random_num}"
        
        bid = frappe.get_doc({
            "doctype": "Bid",
            "bid_code": bid_code,
            "project_name": "New Project",
            "status": "Draft",
            "bid_source": "Standalone"
        })
        bid.insert(ignore_permissions=True)
        frappe.db.commit()
        
        return bid.as_dict()
        
    except Exception as e:
        frappe.log_error(f"Error in create_new_bid: {str(e)}", "Intrakore Estimation")
        return {"error": str(e)}

@frappe.whitelist()
def get_bid_list(filters=None, search=None, page=1, page_size=20):
    """
    Fetch list of bids with filtering, search, and pagination
    """
    try:
        # Build the base query
        conditions = []
        values = {}
        
        # Parse filters if provided as string
        if filters and isinstance(filters, str):
            filters = json.loads(filters)
        
        # Apply status filter
        if filters and filters.get("status"):
            conditions.append("status = %(status)s")
            values["status"] = filters.get("status")
        
        # Apply search filter
        if search:
            conditions.append("""(
                bid_code LIKE %(search)s OR 
                project_name LIKE %(search)s OR 
                client_name LIKE %(search)s
            )""")
            values["search"] = f"%{search}%"
        
        # Build where clause
        where_clause = " AND ".join(conditions) if conditions else "1=1"
        
        # Calculate offset for pagination
        page = int(page) if page else 1
        page_size = int(page_size) if page_size else 20
        offset = (page - 1) * page_size
        
        # Get total count
        total_query = f"""
            SELECT COUNT(*) as total
            FROM `tabBid`
            WHERE {where_clause}
        """
        total_result = frappe.db.sql(total_query, values, as_dict=True)
        total = total_result[0].total if total_result else 0
        
        # Get paginated bids
        query = f"""
            SELECT 
                name,
                bid_code,
                project_name,
                client,
                client_name,
                bid_source,
                status,
                submission_date,
                award_date,
                lead_estimator,
                lead_estimator_name,
                additional_estimators,
                total_priced_value,
                commission_method,
                commission_rate,
                bid_total
            FROM `tabBid`
            WHERE {where_clause}
            ORDER BY creation DESC
            LIMIT %(page_size)s OFFSET %(offset)s
        """
        
        values["page_size"] = page_size
        values["offset"] = offset
        
        bids = frappe.db.sql(query, values, as_dict=True)
        
        # Calculate days left and progress for each bid
        for bid in bids:
            if bid.get("submission_date"):
                due_date = getdate(bid["submission_date"])
                today_date = getdate(today())
                days_left = (due_date - today_date).days
                bid["days_left"] = days_left if days_left > 0 else (0 if days_left == 0 else None)
            else:
                bid["days_left"] = None
            
            progress_map = {
                "Draft": 10, "Mapping": 25, "Tagging": 40, "Pricing": 60,
                "Strategy": 75, "Review": 90, "Submitted": 100,
                "Won": 100, "Lost": 100
            }
            bid["progress"] = progress_map.get(bid.get("status", "Draft"), 0)
        
        total_pages = (total + page_size - 1) // page_size if total > 0 else 1
        
        return {
            "bids": bids,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": total_pages
        }
        
    except Exception as e:
        frappe.log_error(f"Error in get_bid_list: {str(e)}", "Intrakore Estimation API")
        return {"bids": [], "total": 0, "error": str(e)}


# intrakore_estimation/api.py
import frappe
import json
import os
import tempfile
import re
from werkzeug.utils import secure_filename

# Try different import paths
try:
    from intrakore_estimation.utils.boq_parser import BOQParser, UniversalBOQParser, parse_boq_file
except ImportError:
    try:
        from .utils.boq_parser import BOQParser, UniversalBOQParser, parse_boq_file
    except ImportError:
        import sys
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from utils.boq_parser import BOQParser, UniversalBOQParser, parse_boq_file


@frappe.whitelist()
def upload_and_parse_boq():
    """Handle BOQ file upload and parsing - saves original and parsed versions"""
    try:
        import re
        from openpyxl import Workbook
        from io import BytesIO
        
        frappe.log_error("BOQ Upload function called", "Intrakore Estimation")
        
        # FIRST: Ensure Unassigned package exists
        if not frappe.db.exists("Package", "Unassigned"):
            try:
                unassigned = frappe.get_doc({
                    "doctype": "Package",
                    "package_name": "Unassigned",
                    "package_code": "UNAS",
                    "color": "#6b7280",
                    "description": "Unassigned items - need review",
                    "is_active": 1,
                    "default_margin": 0
                })
                unassigned.insert()
                frappe.db.commit()
                print("✅ Created Unassigned package")
            except Exception as e:
                print(f"Error creating Unassigned: {e}")
        
        # Get uploaded file
        if not frappe.request.files:
            return {'error': 'No file uploaded', 'success': False}
        
        file = frappe.request.files.get('file')
        if not file:
            return {'error': 'No file found', 'success': False}
        
        frappe.log_error(f"File received: {file.filename}", "Intrakore Estimation")
        
        # Validate file extension
        filename = secure_filename(file.filename)
        valid_extensions = ['.xlsx', '.xls', '.csv']
        ext = os.path.splitext(filename)[1].lower()
        
        if ext not in valid_extensions:
            return {'error': f'Invalid file type. Please upload {", ".join(valid_extensions)}', 'success': False}
        
        # Read file content
        file_content = file.read()
        
        # ============ SAVE ORIGINAL FILE ============
        # Create a file record for the original BOQ
        original_file = frappe.get_doc({
            "doctype": "File",
            "file_name": f"original_{filename}",
            "content": file_content,
            "is_private": 1,
            "folder": "Home/Attachments"
        })
        original_file.insert()
        frappe.db.commit()
        
        original_file_url = original_file.file_url
        
        # ============ SAVE TEMP FILE FOR PARSING ============
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            tmp.write(file_content)
            temp_path = tmp.name
        
        frappe.log_error(f"Temp file saved at: {temp_path}", "Intrakore Estimation")
        
        # ============ PARSE THE FILE ============
        from intrakore_estimation.utils.boq_parser import BOQParser
        parser = BOQParser()
        parse_result = parser.parse_file(temp_path)
        result = parser.to_json(parse_result)
        
        # ============ CREATE CLEANED EXCEL VERSION ============
        cleaned_wb = Workbook()
        cleaned_ws = cleaned_wb.active
        cleaned_ws.title = "Cleaned BOQ"
        
        # Add headers
        headers = ['Item Ref', 'Description', 'Unit', 'Quantity', 'Package', 'Original Qty', 'Original Unit']
        for col, header in enumerate(headers, 1):
            cleaned_ws.cell(row=1, column=col, value=header)
        
        # Add data rows
        row_num = 2
        for item in result.get('line_items', []):
            cleaned_ws.cell(row=row_num, column=1, value=item.get('item_ref', ''))
            cleaned_ws.cell(row=row_num, column=2, value=item.get('description', ''))
            cleaned_ws.cell(row=row_num, column=3, value=item.get('unit', ''))
            cleaned_ws.cell(row=row_num, column=4, value=item.get('quantity', 0))
            cleaned_ws.cell(row=row_num, column=5, value=item.get('package', 'Unassigned'))
            cleaned_ws.cell(row=row_num, column=6, value=item.get('original_qty', ''))
            cleaned_ws.cell(row=row_num, column=7, value=item.get('original_unit', ''))
            row_num += 1
        
        # Save cleaned file to BytesIO
        cleaned_bytes = BytesIO()
        cleaned_wb.save(cleaned_bytes)
        cleaned_bytes.seek(0)
        
        # Create a file record for the cleaned BOQ
        cleaned_filename = f"cleaned_{filename.replace(ext, '')}_parsed{ext}"
        cleaned_file = frappe.get_doc({
            "doctype": "File",
            "file_name": cleaned_filename,
            "content": cleaned_bytes.getvalue(),
            "is_private": 1,
            "folder": "Home/Attachments"
        })
        cleaned_file.insert()
        frappe.db.commit()
        
        cleaned_file_url = cleaned_file.file_url
        
        # Function to sanitize package name
        def sanitize_package_name(package_name):
            if not package_name or package_name == 'Unassigned':
                return 'Unassigned'
            name = package_name.replace(' ', '-')
            name = name.replace('&', 'and')
            name = re.sub(r'[^a-zA-Z0-9-]', '', name)
            name = re.sub(r'-+', '-', name)
            name = name.strip('-')
            return name if name else 'Unassigned'
        
        # Collect all unique packages from line items
        packages_to_create = set()
        for item in result.get('line_items', []):
            pkg = item.get('package', 'Unassigned')
            if pkg and pkg != 'Unassigned':
                packages_to_create.add(pkg)
        
        for section in result.get('index', {}).get('sections', []):
            pkg = section.get('suggested_package', 'Unassigned')
            if pkg and pkg != 'Unassigned':
                packages_to_create.add(pkg)
        
        # Create each package
        created_packages = []
        for pkg in packages_to_create:
            sanitized = sanitize_package_name(pkg)
            if not frappe.db.exists("Package", sanitized):
                try:
                    new_package = frappe.get_doc({
                        "doctype": "Package",
                        "package_name": sanitized,
                        "package_code": sanitized[:4].upper().replace('-', ''),
                        "color": "#3b82f6",
                        "description": f"{sanitized.replace('-', ' ')} work package",
                        "is_active": 1,
                        "default_margin": 10
                    })
                    new_package.insert()
                    created_packages.append(sanitized)
                    frappe.db.commit()
                except Exception as e:
                    frappe.log_error(f"Error creating package {sanitized}: {str(e)}", "Intrakore Estimation")
        
        # Clean up temp file
        try:
            os.unlink(temp_path)
        except:
            pass
        
        # Add file URLs to result
        result['success'] = True
        result['filename'] = filename
        result['packages_created'] = created_packages
        result['original_boq_url'] = original_file_url
        result['cleaned_boq_url'] = cleaned_file_url
        
        frappe.log_error(
            f"Parse successful: {result.get('total_lines', 0)} line items, "
            f"Created {len(created_packages)} packages, "
            f"Original BOQ saved, Cleaned BOQ saved",
            "Intrakore Estimation"
        )
        
        return result
        
    except Exception as e:
        frappe.log_error(f"BOQ parsing error: {str(e)}", "Intrakore Estimation")
        frappe.log_error(f"Traceback: {frappe.get_traceback()}", "Intrakore Estimation")
        return {'error': str(e), 'success': False}
        
@frappe.whitelist()
def save_parsed_boq(bid_id, boq_data):
    """Save parsed BOQ lines and attach original/cleaned files"""
    try:
        import re
        
        if isinstance(boq_data, str):
            boq_data = json.loads(boq_data)
        
        frappe.log_error(f"Saving BOQ for bid: {bid_id}", "Intrakore Estimation")
        
        # Get the Bid document
        bid = frappe.get_doc("Bid", bid_id)
        
        # Save Index content
        if boq_data.get('index') and boq_data['index'].get('raw_content'):
            bid.client_notes = boq_data['index']['raw_content']
        
        # Attach original BOQ file if URL provided
        if boq_data.get('original_boq_url'):
            # Get the file document
            file_doc = frappe.get_doc("File", {"file_url": boq_data['original_boq_url']})
            if file_doc:
                # Attach to bid
                bid.original_boq = file_doc.file_url
        
        # Attach cleaned BOQ file if URL provided
        if boq_data.get('cleaned_boq_url'):
            file_doc = frappe.get_doc("File", {"file_url": boq_data['cleaned_boq_url']})
            if file_doc:
                bid.priced_boq = file_doc.file_url
        
        # Clear existing BOQ lines
        bid.set('boq_lines', [])
        
        line_count = 0
        
        # Get line items
        line_items = boq_data.get('line_items', []) or boq_data.get('rows', [])
        
        for row in line_items:
            if row.get('type') and row.get('type') != 'line':
                continue
            
            item_ref = row.get('item_ref') or row.get('item') or ''
            description = row.get('description') or ''
            package_name = row.get('package') or row.get('suggested_package') or 'Unassigned'
            unit = row.get('unit') or ''
            qty = row.get('quantity') or row.get('client_qty') or row.get('qty') or 0
            
            if isinstance(qty, str):
                try:
                    numbers = re.findall(r'\d+(?:\.\d+)?', qty)
                    if numbers:
                        qty = float(numbers[0])
                except:
                    qty = 0
            
            bill_name = row.get('bill') or row.get('bill_name') or ''
            section_name = row.get('section') or row.get('section_name') or ''
            warnings = row.get('warnings', [])
            if isinstance(warnings, str):
                warnings = [warnings] if warnings else []
            
            suggested_components = row.get('suggested_components', [])
            suggested_components_json = json.dumps(suggested_components) if suggested_components else ''
            
            bid.append('boq_lines', {
                'item_ref': item_ref[:100] if item_ref else '',
                'description': description[:500] if description else '',
                'unit': unit[:50] if unit else '',
                'client_qty': qty,
                'package': package_name,
                'suggested_package': package_name,
                'suggested_components': suggested_components_json,
                'bill_name': bill_name[:200] if bill_name else '',
                'section_name': section_name[:200] if section_name else '',
                'row_type': 'line',
                'is_confirmed': 0,
                'has_warning': 1 if warnings else 0,
                'warning_message': ', '.join(warnings)[:200] if warnings else '',
                'original_qty': str(row.get('original_qty', '')),
                'original_unit': str(row.get('original_unit', ''))
            })
            line_count += 1
        
        if line_count > 0 and bid.status == 'Draft':
            bid.status = "Mapping"
        
        bid.save(ignore_permissions=True)
        frappe.db.commit()
        
        return {
            "success": True,
            "message": f"Saved {line_count} BOQ lines",
            "total_lines": line_count
        }
        
    except Exception as e:
        frappe.log_error(f"Error saving parsed BOQ: {str(e)}", "Intrakore Estimation")
        frappe.log_error(f"Traceback: {frappe.get_traceback()}", "Intrakore Estimation")
        return {"error": str(e), "success": False}

import re

def sanitize_package_name(package_name):
    """Convert package name to valid format: replace spaces with hyphens, remove special chars"""
    if not package_name:
        return "Unassigned"
    # Replace spaces with hyphens
    name = package_name.replace(' ', '-')
    # Replace & with 'and'
    name = name.replace('&', 'and')
    # Remove any other special characters, keep letters, numbers, hyphens
    name = re.sub(r'[^a-zA-Z0-9-]', '', name)
    # Remove multiple consecutive hyphens
    name = re.sub(r'-+', '-', name)
    # Remove leading/trailing hyphens
    name = name.strip('-')
    return name if name else "Unassigned"
